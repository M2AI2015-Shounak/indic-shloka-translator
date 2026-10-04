"""End-to-end pipeline with a mock backend (offline, no GPU)."""

from __future__ import annotations

import json
import logging

from openpyxl import load_workbook

from shloka_translator.output.excel_writer import SHEET_ORDER, TRANSLATION_COLUMNS
from shloka_translator.pipeline.orchestrator import RunOptions, attach_file_log, detach_file_log, run_pipeline
from tests.helpers import FIXTURES, REFERENCE, MockBackend, make_config


def _run(cfg, backend, **kw):
    """Run the pipeline with a mock backend.

    Args:
        cfg: Config.
        backend: Mock backend.
        **kw: RunOptions overrides.

    Returns:
        RunResult.
    """
    opts = RunOptions(
        model="mock-model", backend_factory=lambda _n: backend, use_widgets=False, openai_key="", hf_token="", **kw
    )
    return run_pipeline(cfg, opts)


def _rows(path):
    """Read the Translations sheet as dicts.

    Args:
        path: Workbook path.

    Returns:
        List of row dicts.
    """
    ws = load_workbook(path)["Translations"]
    values = list(ws.iter_rows(values_only=True))
    return [dict(zip(values[0], v, strict=True)) for v in values[1:]]


def test_smoke_five_records(tmp_path) -> None:
    """5-record run: correctly named workbook with all six sheets; Verse kept as text.

    Args:
        tmp_path: Temp dir.
    """
    cfg = make_config(tmp_path, [REFERENCE])
    res = _run(cfg, MockBackend(), max_records=5)
    assert res.ok
    m = res.models[0]
    assert m.output_path.name.startswith("mock-model_") and m.output_path.suffix == ".xlsx"
    assert m.log_path.name == m.output_path.stem + ".log"
    wb = load_workbook(m.output_path)
    assert wb.sheetnames == SHEET_ORDER
    rows = _rows(m.output_path)
    assert len(rows) == 5 and list(rows[0]) == TRANSLATION_COLUMNS
    assert all(r["Processing_Status"] == "Success" for r in rows)
    assert rows[1]["Remarks"].endswith("Note: ambiguous word")
    assert rows[0]["Verse"] == "1.1" and isinstance(rows[0]["Verse"], str)
    assert rows[0]["Total_Tokens"] == 18 and rows[0]["Detected_Language"] == "Sanskrit"
    summary = dict(load_workbook(m.output_path)["Run_Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Records total"] == 5 and summary["Success"] == 5


def test_skipped_failed_and_rejected(tmp_path) -> None:
    """Unsupported languages are Skipped, one failing record is isolated, a bad file is rejected.

    Args:
        tmp_path: Temp dir.
    """
    cfg = make_config(tmp_path, [FIXTURES / "mixed_languages.csv", FIXTURES / "invalid_extra_column.csv"])
    res = _run(cfg, MockBackend(fail_ids={"M01"}))
    rows = {r["ID"]: r for r in _rows(res.models[0].output_path)}
    assert rows["M27"]["Processing_Status"] == "Skipped" and "Unsupported language: English" in rows["M27"]["Remarks"]
    assert rows["M28"]["Processing_Status"] == "Skipped" and "Urdu" in rows["M28"]["Remarks"]
    assert rows["M31"]["Processing_Status"] == "Skipped"
    assert rows["M01"]["Processing_Status"] == "Failed" and "simulated failure" in rows["M01"]["Remarks"]
    assert rows["M02"]["Processing_Status"] == "Success"
    val = list(load_workbook(res.models[0].output_path)["Input_Validation"].iter_rows(values_only=True))
    rejected = [v for v in val if v[2] == "Rejected"]
    assert rejected and rejected[0][0] == "invalid_extra_column.csv" and "extra=['Notes']" in rejected[0][4]


def test_no_valid_input_stops_gracefully(tmp_path) -> None:
    """If no file passes validation the run returns a message instead of raising.

    Args:
        tmp_path: Temp dir.
    """
    cfg = make_config(tmp_path, [FIXTURES / "invalid_missing_column.csv"])
    res = _run(cfg, MockBackend())
    assert not res.ok and "No input file passed validation" in res.message


def test_resume_from_checkpoint(tmp_path) -> None:
    """Interrupt mid-run → partial output; re-run resumes without re-calling finished records.

    Args:
        tmp_path: Temp dir.
    """
    cfg = make_config(tmp_path, [REFERENCE], concurrency={"openai_max_workers": 1})
    first = MockBackend(interrupt_after=4)
    res1 = _run(cfg, first, max_records=10)
    m1 = res1.models[0]
    assert m1.interrupted and m1.output_path.exists()
    statuses = [r["Processing_Status"] for r in _rows(m1.output_path)]
    assert statuses.count("Success") == 4 and statuses.count("Failed") == 6
    second = MockBackend()
    res2 = _run(cfg, second, max_records=10)
    assert second.calls == 6, "finished records must not be re-billed"
    assert res2.models[0].counts["Cached"] == 4 and res2.models[0].counts["Success"] == 10
    third = MockBackend()
    _run(cfg, third, max_records=10, resume=False)
    assert third.calls == 10


def test_no_secrets_in_outputs_or_logs(tmp_path, monkeypatch) -> None:
    """Secrets never reach logs or workbooks.

    Args:
        tmp_path: Temp dir.
        monkeypatch: Pytest monkeypatch.
    """
    fake_key = "sk-test" + "A1b2C3d4E5f6G7h8I9j0"
    monkeypatch.setenv("OPENAI_API_KEY", fake_key)
    cfg = make_config(tmp_path, [REFERENCE])
    opts = RunOptions(model="mock-model", backend_factory=lambda _n: MockBackend(), use_widgets=False, max_records=2)
    res = run_pipeline(cfg, opts)
    probe = tmp_path / "logs" / "probe.log"
    handler = attach_file_log(probe)
    try:
        logging.getLogger("x").warning("leak attempt %s and hf_%s", fake_key, "abcdefghijklmnop")
    finally:
        detach_file_log(handler)
    assert "[REDACTED]" in probe.read_text(encoding="utf-8")
    for path in [*(tmp_path / "logs").glob("*.log"), res.models[0].output_path]:
        if path.suffix == ".log":
            assert fake_key not in path.read_text(encoding="utf-8")
    wb = load_workbook(res.models[0].output_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and fake_key in v for v in row)
    for jf in (tmp_path / "cache").rglob("*.jsonl"):
        assert fake_key not in jf.read_text(encoding="utf-8")
    assert json.dumps(cfg.snapshot()).find(fake_key) == -1
