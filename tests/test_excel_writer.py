"""Workbook layout: sheets, column order, colours, sanitisation, truncation."""

from __future__ import annotations

from openpyxl import load_workbook

from shloka_translator.monitoring.gpu_monitor import GPUStatic, RunStats, Samples, build_metrics_rows
from shloka_translator.monitoring.gpu_specs import lookup_gpu_spec
from shloka_translator.output.excel_writer import (
    SHEET_ORDER,
    TRANSLATION_COLUMNS,
    WorkbookData,
    write_comparison,
    write_workbook,
)


def _row(i: int, status: str, meaning: str = "m") -> dict:
    """Build a translation row.

    Args:
        i: ID.
        status: Status.
        meaning: Translation.

    Returns:
        Row dict.
    """
    return {
        "ID": str(i),
        "Source": "S",
        "Verse": "1.10",
        "Shloka": "राम\x07",
        "Detected_Language": "Sanskrit",
        "Generated_Meaning": meaning,
        "Word_Segmentation": "a | b",
        "Word_Meanings": "a: x\nb: y",
        "Input_Tokens": 1,
        "Output_Tokens": 2,
        "Total_Tokens": 3,
        "Processing_Status": status,
        "Remarks": "",
    }


def test_workbook_layout(tmp_path) -> None:
    """Six sheets in order, exact columns, status colours, illegal chars removed, long cells truncated.

    Args:
        tmp_path: Temp dir.
    """
    rows = [_row(1, "Success"), _row(2, "Skipped"), _row(3, "Failed", "x" * 40000)]
    st = GPUStatic(
        name="Tesla T4",
        total_mem_gb=15.0,
        max_graphics_clock=1590,
        max_mem_clock=5001,
        bus_width=256,
        pcie_gen_max=3,
        pcie_width_max=16,
    )
    sm = Samples(
        t=[0, 1, 2],
        util=[50, 60, 70],
        mem_util=[10, 20, 30],
        mem_used_gb=[1, 2, 3],
        power_w=[30, 40, 50],
        temp_c=[40, 41, 42],
        gfx_clock=[1500, 1500, 1500],
        mem_clock=[5000, 5000, 5000],
    )
    gpu_rows = build_metrics_rows(
        st, sm, lookup_gpu_spec("Tesla T4"), RunStats(3, 30, 60, [0.1, 0.2, 0.3], 2.0, "16", "hf_seq2seq"), 0.02
    )
    assert len(gpu_rows) == 20
    path = write_workbook(tmp_path / "t.xlsx", WorkbookData(rows, gpu_rows, "", [("Model", "x")], [], [], []))
    wb = load_workbook(path)
    assert wb.sheetnames == SHEET_ORDER
    ws = wb["Translations"]
    assert [c.value for c in ws[1]] == TRANSLATION_COLUMNS
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws["D2"].value == "राम"
    assert ws["C2"].value == "1.10"
    status_col = TRANSLATION_COLUMNS.index("Processing_Status") + 1
    colours = [ws.cell(row=r, column=status_col).fill.fgColor.rgb[-6:] for r in (2, 3, 4)]
    assert colours == ["C6EFCE", "FFEB9C", "FFC7CE"]
    assert len(ws["F4"].value) <= 32767
    assert "truncated" in ws.cell(row=4, column=len(TRANSLATION_COLUMNS)).value
    params = [r[1] for r in wb["GPU_Metrics"].iter_rows(min_row=2, values_only=True)]
    assert params[0] == "GPU Utilization" and params[-1] == "Compute Occupancy / SM Utilization"
    bw = next(r for r in gpu_rows if r[1] == "Memory Bandwidth")
    assert bw[2] == "320.1" and bw[4] == "Derived"


def test_comparison(tmp_path) -> None:
    """Comparison workbook has one meaning column per model.

    Args:
        tmp_path: Temp dir.
    """
    rows = [_row(1, "Success", "A"), _row(2, "Failed", "")]
    path = write_comparison(tmp_path / "c.xlsx", rows, {"m1": rows, "m2": rows})
    ws = load_workbook(path)["Comparison"]
    header = [c.value for c in ws[1]]
    assert header[-2:] == ["m1 — Generated_Meaning", "m2 — Generated_Meaning"]
    assert ws.cell(row=3, column=6).value == "[Failed]"
