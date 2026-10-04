"""Read CSV / XLSX / JSON / JSONL input files into all-text DataFrames.

Every cell is read as text (never coerced to float), NFC-normalised and
stripped of zero-width/control characters. Each XLSX sheet becomes its own
:class:`LoadedTable` tagged ``file::sheet`` so it can be validated on its own.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

from shloka_translator.utils.text import normalize_text, verse_to_text

logger = logging.getLogger(__name__)

CSV_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "utf-16")
CSV_DELIMITERS = ",;\t|"
IGNORED_NAMES = {"readme.md"}


@dataclass
class LoadedTable:
    """One table read from an input file (a CSV/JSON file or one XLSX sheet).

    Attributes:
        source_name: Human-readable origin, ``"file.xlsx::Sheet1"`` or ``"file.csv"``.
        file_name: File name only.
        sheet: Sheet name for XLSX, else ``""``.
        frame: All-text DataFrame with the file's own header (``None`` on error).
        error: Reason the file could not be read, else ``None``.
        numeric_verse_rows: 0-based row positions whose ``Verse`` cell was a
            non-integral number in the XLSX (e.g. ``1.1``), which may have lost
            a trailing zero (``1.10`` → ``1.1``) before it ever reached us.
    """

    source_name: str
    file_name: str
    sheet: str = ""
    frame: pd.DataFrame | None = None
    error: str | None = None
    numeric_verse_rows: set[int] = field(default_factory=set)


def cell_to_text(value: Any) -> str:
    """Convert one raw cell to clean text without numeric corruption.

    Args:
        value: Raw value from openpyxl/JSON (int, float, str, None, ...).

    Returns:
        Clean text; integral numbers have no trailing ``.0``.

    Example:
        >>> cell_to_text(5.0), cell_to_text(None), cell_to_text(" a ")
        ('5', '', 'a')
    """
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return verse_to_text(value)
    return normalize_text(value)


def discover_files(folder: Path, extensions: list[str], recursive: bool = False) -> list[Path]:
    """List candidate input files in a deterministic order.

    Hidden files, Excel lock files (``~$...``) and ``README.md`` are ignored.

    Args:
        folder: Input folder.
        extensions: Accepted lowercase extensions, e.g. ``[".csv", ".xlsx"]``.
        recursive: Search sub-folders too.

    Returns:
        Sorted list of file paths.

    Raises:
        FileNotFoundError: If ``folder`` does not exist.
    """
    if not folder.is_dir():
        raise FileNotFoundError(f"Input folder not found: {folder}")
    pattern = "**/*" if recursive else "*"
    exts = {e.lower() for e in extensions}
    files = []
    for path in folder.glob(pattern):
        name = path.name
        if not path.is_file() or name.startswith(".") or name.startswith("~$"):
            continue
        if name.lower() in IGNORED_NAMES or path.suffix.lower() not in exts:
            continue
        if any(part.startswith(".") for part in path.relative_to(folder).parts[:-1]):
            continue
        files.append(path)
    return sorted(files, key=lambda p: str(p.relative_to(folder)).lower())


def _frame_from_rows(header: list[Any], rows: list[list[Any]]) -> pd.DataFrame:
    """Build an all-text DataFrame from a header row and data rows.

    Args:
        header: Raw header cells.
        rows: Raw data rows (may be ragged).

    Returns:
        DataFrame whose column names are the header as text (duplicates kept).
    """
    columns = [cell_to_text(h) if h is not None else "" for h in header]
    width = len(columns)
    data = []
    for row in rows:
        cells = [cell_to_text(v) for v in list(row)[:width]]
        cells += [""] * (width - len(cells))
        data.append(cells)
    frame = pd.DataFrame(data, columns=pd.Index(columns, dtype=object), dtype=object)
    return frame


def _drop_blank_rows(frame: pd.DataFrame) -> pd.DataFrame:
    """Remove rows whose cells are all empty strings.

    Args:
        frame: All-text frame.

    Returns:
        Frame without fully blank rows, index reset.
    """
    if frame.empty:
        return frame
    mask = frame.apply(lambda r: any(str(v).strip() for v in r.values), axis=1)
    return frame[mask].reset_index(drop=True)


def read_xlsx(path: Path) -> list[LoadedTable]:
    """Read every non-empty sheet of an XLSX workbook as text.

    Args:
        path: Workbook path.

    Returns:
        One :class:`LoadedTable` per non-empty sheet (or one error entry).
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        return [LoadedTable(path.name, path.name, error=f"Unreadable workbook: {type(exc).__name__}: {exc}")]
    tables: list[LoadedTable] = []
    try:
        for ws in wb.worksheets:
            raw_rows = [list(r) for r in ws.iter_rows(values_only=True)]
            raw_rows = [r for r in raw_rows if any(v is not None and str(v).strip() for v in r)]
            if not raw_rows:
                logger.info("Skipping empty sheet %s::%s", path.name, ws.title)
                continue
            header, body = raw_rows[0], raw_rows[1:]
            while header and (header[-1] is None or str(header[-1]).strip() == ""):
                header = header[:-1]
            frame = _frame_from_rows(header, body)
            numeric_rows: set[int] = set()
            header_text = [cell_to_text(h).strip() for h in header]
            if "Verse" in header_text:
                vidx = header_text.index("Verse")
                for i, row in enumerate(body):
                    v = row[vidx] if vidx < len(row) else None
                    if isinstance(v, float) and not v.is_integer():
                        numeric_rows.add(i)
            source = f"{path.name}::{ws.title}"
            tables.append(LoadedTable(source, path.name, ws.title, _drop_blank_rows(frame), None, numeric_rows))
    finally:
        wb.close()
    if not tables:
        return [LoadedTable(path.name, path.name, error="Empty workbook (no non-empty sheets)")]
    return tables


def read_csv(path: Path) -> LoadedTable:
    """Read a CSV file as text, trying several encodings and delimiters.

    Args:
        path: CSV path.

    Returns:
        A :class:`LoadedTable` (with ``error`` set on failure).
    """
    raw = path.read_bytes()
    if not raw.strip():
        return LoadedTable(path.name, path.name, error="Empty file")
    last_error = ""
    for encoding in CSV_ENCODINGS:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            last_error = f"decode {encoding}: {exc}"
            continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=CSV_DELIMITERS)
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        try:
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            rows = [r for r in reader]
        except csv.Error as exc:
            last_error = f"parse: {exc}"
            continue
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            return LoadedTable(path.name, path.name, error="Empty file")
        frame = _frame_from_rows(rows[0], rows[1:])
        logger.debug("Read %s with encoding=%s delimiter=%r", path.name, encoding, delimiter)
        return LoadedTable(path.name, path.name, frame=_drop_blank_rows(frame))
    return LoadedTable(path.name, path.name, error=f"Could not decode CSV ({last_error})")


def _records_to_frame(records: list[Any]) -> pd.DataFrame:
    """Convert a list of JSON objects to an all-text frame.

    Args:
        records: Parsed JSON records (must all be objects).

    Returns:
        Frame with columns in first-seen order.

    Raises:
        ValueError: If any record is not a JSON object.
    """
    columns: list[str] = []
    for rec in records:
        if not isinstance(rec, dict):
            raise ValueError("every record must be a JSON object")
        for key in rec:
            if key not in columns:
                columns.append(key)
    rows = [[rec.get(c) for c in columns] for rec in records]
    return _frame_from_rows(columns, rows)


def read_json(path: Path) -> LoadedTable:
    """Read JSON (list of records or ``{"records": [...]}``) or JSON Lines.

    Args:
        path: ``.json`` or ``.jsonl`` file.

    Returns:
        A :class:`LoadedTable`; unsupported shapes produce an ``error``.
    """
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    if not text.strip():
        return LoadedTable(path.name, path.name, error="Empty file")
    records: list[Any] | None = None
    try:
        data = json.loads(text)
        if isinstance(data, list):
            records = data
        elif isinstance(data, dict) and isinstance(data.get("records"), list):
            records = data["records"]
        else:
            return LoadedTable(
                path.name,
                path.name,
                error='Unsupported JSON shape (expected a list of records, {"records": [...]}, or JSON Lines)',
            )
    except json.JSONDecodeError:
        try:
            records = [json.loads(line) for line in text.splitlines() if line.strip()]
        except json.JSONDecodeError as exc:
            return LoadedTable(path.name, path.name, error=f"Invalid JSON / JSON Lines: {exc}")
    if not records:
        return LoadedTable(path.name, path.name, error="Empty file (no records)")
    try:
        frame = _records_to_frame(records)
    except ValueError as exc:
        return LoadedTable(path.name, path.name, error=f"Unsupported JSON shape ({exc})")
    return LoadedTable(path.name, path.name, frame=_drop_blank_rows(frame))


def load_file(path: Path) -> list[LoadedTable]:
    """Dispatch to the right reader based on the file extension.

    Args:
        path: Input file.

    Returns:
        List of loaded tables (XLSX may yield several).
    """
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            return read_xlsx(path)
        if suffix == ".csv":
            return [read_csv(path)]
        if suffix in (".json", ".jsonl"):
            return [read_json(path)]
    except Exception as exc:
        logger.exception("Failed to read %s", path.name)
        return [LoadedTable(path.name, path.name, error=f"{type(exc).__name__}: {exc}")]
    return [LoadedTable(path.name, path.name, error=f"Unsupported extension {suffix}")]


def load_folder(folder: Path, extensions: list[str], recursive: bool = False) -> list[LoadedTable]:
    """Load every supported file in the input folder.

    Args:
        folder: Input folder path.
        extensions: Accepted extensions.
        recursive: Whether to search sub-folders.

    Returns:
        Loaded tables in deterministic (file, sheet) order.
    """
    tables: list[LoadedTable] = []
    for path in discover_files(folder, extensions, recursive):
        loaded = load_file(path)
        if recursive:
            rel = str(path.relative_to(folder))
            for t in loaded:
                t.file_name = rel
                t.source_name = f"{rel}::{t.sheet}" if t.sheet else rel
        tables.extend(loaded)
        logger.info("Loaded %s → %d table(s)", path.name, len(loaded))
    return tables
