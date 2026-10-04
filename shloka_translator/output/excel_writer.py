"""Write the six-sheet output workbook (and the ALL-MODELS comparison) with openpyxl."""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from shloka_translator.utils.text import excel_safe

logger = logging.getLogger(__name__)

TRANSLATION_COLUMNS: list[str] = [
    "ID",
    "Source",
    "Verse",
    "Shloka",
    "Detected_Language",
    "Generated_Meaning",
    "Word_Segmentation",
    "Word_Meanings",
    "Input_Tokens",
    "Output_Tokens",
    "Total_Tokens",
    "Processing_Status",
    "Remarks",
]
GPU_COLUMNS = ["#", "Parameter", "Value", "Unit", "Method", "Notes"]
VALIDATION_COLUMNS = ["File", "Sheet", "Status", "Rows", "Reason"]
DETECTION_COLUMNS = [
    "ID",
    "Detected_Language",
    "Language_Code",
    "Script",
    "Confidence",
    "Method",
    "Source_File",
    "Source_Sheet",
    "Row_Number",
]
TOKEN_COLUMNS = [
    "ID",
    "Translation_Input_Tokens",
    "Translation_Output_Tokens",
    "Translation_Total_Tokens",
    "Reasoning_Tokens",
    "Helper_Input_Tokens",
    "Helper_Output_Tokens",
    "LID_Tiebreaker_Input_Tokens",
    "LID_Tiebreaker_Output_Tokens",
    "Retries",
    "Cached",
]
SHEET_ORDER = ["Translations", "GPU_Metrics", "Run_Summary", "Input_Validation", "Detection_Audit", "Token_Breakdown"]
STATUS_FILLS = {"Success": "C6EFCE", "Skipped": "FFEB9C", "Failed": "FFC7CE"}
WIDTHS = {
    "ID": 10,
    "Source": 22,
    "Verse": 10,
    "Shloka": 60,
    "Detected_Language": 14,
    "Generated_Meaning": 70,
    "Word_Segmentation": 50,
    "Word_Meanings": 60,
    "Input_Tokens": 12,
    "Output_Tokens": 13,
    "Total_Tokens": 12,
    "Processing_Status": 16,
    "Remarks": 60,
}
WRAP_COLUMNS = {
    "Shloka",
    "Generated_Meaning",
    "Word_Segmentation",
    "Word_Meanings",
    "Remarks",
    "Value",
    "Notes",
    "Reason",
    "Method",
}


@dataclass
class WorkbookData:
    """Everything that goes into one model's workbook.

    Attributes:
        translations: Rows for ``Translations`` (dicts keyed by :data:`TRANSLATION_COLUMNS`).
        gpu_rows: Rows for ``GPU_Metrics``.
        gpu_note: Optional note placed above the GPU table.
        summary: ``(key, value)`` pairs for ``Run_Summary``.
        validation: Rows for ``Input_Validation``.
        detection: Rows for ``Detection_Audit``.
        tokens: Rows for ``Token_Breakdown``.
    """

    translations: list[dict[str, Any]] = field(default_factory=list)
    gpu_rows: list[list[Any]] = field(default_factory=list)
    gpu_note: str = ""
    summary: list[tuple[str, Any]] = field(default_factory=list)
    validation: list[list[Any]] = field(default_factory=list)
    detection: list[list[Any]] = field(default_factory=list)
    tokens: list[list[Any]] = field(default_factory=list)


def _style_header(ws: Any, ncols: int, header_row: int = 1) -> None:
    """Bold header, freeze panes below it, enable autofilter.

    Args:
        ws: Worksheet.
        ncols: Number of columns.
        header_row: Row index of the header.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", fgColor="DDEBF7")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if ws.max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{ws.max_row}"


def _write_table(
    ws: Any, columns: list[str], rows: list[list[Any]], start_row: int = 1, widths: dict[str, int] | None = None
) -> list[int]:
    """Write a header and rows with wrapping and widths.

    Args:
        ws: Worksheet.
        columns: Header names.
        rows: Data rows.
        start_row: Header row index.
        widths: Optional column → width map.

    Returns:
        Indices (0-based) of rows that had a cell truncated.
    """
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    truncated_rows: list[int] = []
    for c, name in enumerate(columns, 1):
        ws.cell(row=start_row, column=c, value=name)
    for r, row in enumerate(rows):
        for c, value in enumerate(row, 1):
            clean, cut = excel_safe(value)
            if cut:
                truncated_rows.append(r)
            cell = ws.cell(row=start_row + 1 + r, column=c, value=clean)
            if columns[c - 1] in WRAP_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
    for c, name in enumerate(columns, 1):
        width = (widths or {}).get(name)
        if width is None:
            longest = max([len(str(name))] + [len(str(row[c - 1])) for row in rows[:200] if c - 1 < len(row)])
            width = min(60, max(10, longest + 2))
        ws.column_dimensions[get_column_letter(c)].width = width
    _style_header(ws, len(columns), start_row)
    return truncated_rows


def _translations_sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
    """Fill the ``Translations`` sheet (with status colours and truncation notes).

    Args:
        ws: Worksheet.
        rows: Translation row dicts.
    """
    from openpyxl.styles import PatternFill

    matrix = [[row.get(c, "") for c in TRANSLATION_COLUMNS] for row in rows]
    remarks_idx = TRANSLATION_COLUMNS.index("Remarks")
    for row in matrix:
        for i, value in enumerate(row):
            if isinstance(value, str) and len(value) > 32_767 and i != remarks_idx:
                note = f"{TRANSLATION_COLUMNS[i]} truncated to Excel's 32,767-character cell limit"
                row[remarks_idx] = f"{row[remarks_idx]}; {note}".strip("; ")
    _write_table(ws, TRANSLATION_COLUMNS, matrix, widths=WIDTHS)
    status_col = TRANSLATION_COLUMNS.index("Processing_Status") + 1
    for r in range(2, len(matrix) + 2):
        cell = ws.cell(row=r, column=status_col)
        color = STATUS_FILLS.get(str(cell.value))
        if color:
            cell.fill = PatternFill("solid", fgColor=color)


def write_workbook(path: Path, data: WorkbookData) -> Path:
    """Write the model workbook with the six required sheets.

    Args:
        path: Destination ``.xlsx`` (must not exist; see :func:`naming.unique_path`).
        data: Sheet contents.

    Returns:
        The written path.

    Raises:
        OSError: If the file cannot be written.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Translations"
    _translations_sheet(ws, data.translations)

    ws = wb.create_sheet("GPU_Metrics")
    start = 1
    if data.gpu_note:
        ws.cell(row=1, column=1, value=data.gpu_note).font = Font(italic=True, bold=True)
        start = 3
    _write_table(
        ws,
        GPU_COLUMNS,
        data.gpu_rows,
        start_row=start,
        widths={"#": 5, "Parameter": 34, "Value": 60, "Unit": 22, "Method": 40, "Notes": 70},
    )

    ws = wb.create_sheet("Run_Summary")
    _write_table(
        ws,
        ["Item", "Value"],
        [[k, v if isinstance(v, (int, float)) else str(v)] for k, v in data.summary],
        widths={"Item": 34, "Value": 110},
    )

    ws = wb.create_sheet("Input_Validation")
    _write_table(
        ws,
        VALIDATION_COLUMNS,
        data.validation,
        widths={"File": 36, "Sheet": 16, "Status": 12, "Rows": 8, "Reason": 100},
    )
    ws = wb.create_sheet("Detection_Audit")
    _write_table(ws, DETECTION_COLUMNS, data.detection)
    ws = wb.create_sheet("Token_Breakdown")
    _write_table(ws, TOKEN_COLUMNS, data.tokens)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("Wrote %s", path)
    return path


def write_comparison(path: Path, base_rows: list[dict[str, Any]], per_model: dict[str, list[dict[str, Any]]]) -> Path:
    """Write the ALL-MODELS side-by-side comparison workbook.

    Args:
        path: Destination ``.xlsx``.
        base_rows: Translation rows of any model (for ID/Source/Verse/Shloka/language).
        per_model: Model name → its translation rows (same order).

    Returns:
        The written path.
    """
    from openpyxl import Workbook

    models = list(per_model)
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    cols = ["ID", "Source", "Verse", "Shloka", "Detected_Language"] + [f"{m} — Generated_Meaning" for m in models]
    rows = []
    for i, base in enumerate(base_rows):
        row = [base.get(c, "") for c in cols[:5]]
        for m in models:
            mr = per_model[m][i] if i < len(per_model[m]) else {}
            row.append(mr.get("Generated_Meaning", "") or f"[{mr.get('Processing_Status', 'N/A')}]")
        rows.append(row)
    widths = {c: 50 for c in cols[5:]} | {"Shloka": 50}
    for c in cols[5:]:
        WRAP_COLUMNS.add(c)
    _write_table(ws, cols, rows, widths=widths)
    ws = wb.create_sheet("Status_Comparison")
    scols = ["ID"] + [f"{m} — Status" for m in models] + [f"{m} — Total_Tokens" for m in models]
    srows = []
    for i, base in enumerate(base_rows):
        r = [base.get("ID", "")]
        r += [per_model[m][i].get("Processing_Status", "") if i < len(per_model[m]) else "" for m in models]
        r += [per_model[m][i].get("Total_Tokens", 0) if i < len(per_model[m]) else 0 for m in models]
        srows.append(r)
    _write_table(ws, scols, srows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("Wrote comparison %s", path)
    return path
